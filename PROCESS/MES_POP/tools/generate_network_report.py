import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def create_excel_report():
    wb = openpyxl.Workbook()
    
    font_family = "Calibri"
    
    title_font = Font(name=font_family, size=15, bold=True, color="FFFFFF")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="E0E0E0")
    section_font = Font(name=font_family, size=12, bold=True, color="1B365D")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=10, bold=True, color="000000")
    regular_font = Font(name=font_family, size=10, color="000000")
    kpi_title_font = Font(name=font_family, size=9, bold=True, color="595959")
    
    # Fills
    primary_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Deep Navy
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")  # Steel Blue
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    # Status Fills
    done_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")    # Soft Green
    done_font = Font(name=font_family, size=10, bold=True, color="375623")
    
    pending_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft Yellow
    pending_font = Font(name=font_family, size=10, bold=True, color="833C0C")
    
    blocked_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft Red/Orange
    blocked_font = Font(name=font_family, size=10, bold=True, color="C65911")
    
    # Borders
    thin_gray = Side(style='thin', color='D9D9D9')
    med_navy = Side(style='medium', color='1B365D')
    double_bottom = Side(style='double', color='1B365D')
    
    cell_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    total_border = Border(top=thin_gray, bottom=double_bottom, left=thin_gray, right=thin_gray)
    
    # Alignments
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')

    # =============================================================
    # SHEET 1: TIẾN ĐỘ KÉO CÁP CELL LINE
    # =============================================================
    ws1 = wb.active
    ws1.title = "1. TIEN DO KEO CAP CELL LINE"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws1.merge_cells("A1:J1")
    ws1["A1"] = "VINATECH CO., LTD. — IT INFRASTRUCTURE & FACTORY NETWORK REPORT"
    ws1["A1"].font = Font(name=font_family, size=10, bold=True, color="D9E1F2")
    ws1["A1"].fill = primary_fill
    ws1["A1"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws1.row_dimensions[1].height = 20

    ws1.merge_cells("A2:J2")
    ws1["A2"] = "BÁO CÁO TIẾN ĐỘ HẠ TẦNG DÂY MẠNG CELL LINE (CELL 라인 네트워크 케이블 포설 현황)"
    ws1["A2"].font = title_font
    ws1["A2"].fill = primary_fill
    ws1["A2"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws1.row_dimensions[2].height = 28

    ws1.merge_cells("A3:J3")
    ws1["A3"] = "Ngày khảo sát: 07/10/2026 | Người thực hiện: Nguyễn Văn Đức (vanduc - EA Team) | Phụ trách / Báo cáo: Mr. Hải / Mr. Lee Sang Hwan"
    ws1["A3"].font = subtitle_font
    ws1["A3"].fill = primary_fill
    ws1["A3"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws1.row_dimensions[3].height = 18

    # KPI Summary Cards
    kpi_configs = [
        ("B5:C5", "B6:C6", "TỔNG DÂY CẦN THIẾT (총 소요)", "30 Dây (Ports)", "B5", "B6", primary_fill, Font(name=font_family, size=15, bold=True, color="1B365D")),
        ("D5:E5", "D6:E6", "ĐÃ KÉO XONG (포설 완료)", "27 Dây (90.0%)", "D5", "D6", done_fill, Font(name=font_family, size=15, bold=True, color="375623")),
        ("F5:G5", "F6:G6", "CẦN KÉO THÊM (진행 가능/잔여)", "3 Dây (Aging L6, Winding L3)", "F5", "F6", pending_fill, Font(name=font_family, size=12, bold=True, color="833C0C")),
        ("H5:I5", "H6:I6", "NGHẼN HẠ TẦNG (인프라 장애/대기)", "14 Vị trí (Winding 6, Ass'y 8)", "H5", "H6", blocked_fill, Font(name=font_family, size=12, bold=True, color="C65911")),
    ]
    
    ws1.row_dimensions[5].height = 18
    ws1.row_dimensions[6].height = 26
    
    for t_range, v_range, title, val, t_cell, v_cell, fill, f_font in kpi_configs:
        ws1.merge_cells(t_range)
        ws1[t_cell] = title
        ws1[t_cell].font = kpi_title_font
        ws1[t_cell].fill = gray_fill
        ws1[t_cell].alignment = center_align
        ws1[t_cell].border = cell_border
        
        ws1.merge_cells(v_range)
        ws1[v_cell] = val
        ws1[v_cell].font = f_font
        ws1[v_cell].fill = fill
        ws1[v_cell].alignment = center_align
        ws1[v_cell].border = cell_border

    # Section 1: Summary Table
    ws1.cell(row=8, column=1, value="I. BẢNG TỔNG HỢP THEO KHU VỰC CÔNG ĐOẠN (공정 구역별 요약)").font = section_font
    
    summary_headers = ["STT", "Khu Vực Công Đoạn (공정 구역)", "Tổng Dây Khả Thi", "Đã Kéo Xong", "Cần Kéo Thêm", "Vị Trí Vướng Hạ Tầng", "Tỷ Lệ Hoàn Thành (%)", "Đánh Giá / Hiện Trạng Kỹ Thuật"]
    ws1.row_dimensions[9].height = 24
    for col_idx, h in enumerate(summary_headers, start=1):
        c = ws1.cell(row=9, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = cell_border
        
    summary_data = [
        (1, "Aging (노화공정 - Hồi/Ủ điện tích)", 6, 5, 1, 0, "=D10/C10", "🟢 5/5 máy 35105 đã xong. Còn thiếu 1 dây cho cụm 2245 Aging #1."),
        (2, "Winding (권취공정 - Quấn cuộn)", 12, 10, 2, 6, "=D11/C11", "🟡 L1(2/2), L2(6/6) xong. L3 còn thiếu 2 dây. L4-9 bị vướng máng cáp."),
        (3, "Assembly (조립공정 - Lắp ráp Cell)", 10, 10, 0, 8, "=D12/C12", "🟢 Cụm máy 35105 (Line 9, 10) đã kéo đủ 10/10 dây. Line 1-8 vướng hạ tầng."),
        (4, "Phụ trợ: Partleader/VPSX & X-Ray", 2, 2, 0, 0, "=D13/C13", "🟢 Đã xong 1 Uplink vào Switch 8P Partleader (chia máy bàn) & 1 dây X-Ray.")
    ]
    
    for r_idx, row_values in enumerate(summary_data, start=10):
        ws1.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if c_idx in [1, 3, 4, 5, 6]:
                cell.alignment = center_align
            elif c_idx == 7:
                cell.alignment = center_align
                cell.number_format = '0.0%'
                cell.font = bold_font
            else:
                cell.alignment = left_align
                
    # Total row
    ws1.row_dimensions[14].height = 22
    ws1.cell(row=14, column=1, value="").border = total_border
    c_tot_label = ws1.cell(row=14, column=2, value="TỔNG CỘNG KHU VỰC (합계)")
    c_tot_label.font = bold_font
    c_tot_label.alignment = left_align
    c_tot_label.border = total_border
    
    for c_idx, col_let in [(3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')]:
        c = ws1.cell(row=14, column=c_idx, value=f"=SUM({col_let}10:{col_let}13)")
        c.font = bold_font
        c.alignment = center_align
        c.border = total_border
        
    c_tot_rate = ws1.cell(row=14, column=7, value="=D14/C14")
    c_tot_rate.font = bold_font
    c_tot_rate.alignment = center_align
    c_tot_rate.number_format = '0.0%'
    c_tot_rate.border = total_border
    
    c_tot_rem = ws1.cell(row=14, column=8, value="Tổng 30 dây khả thi: Đã xong 27 dây (90.0%). Cần triển khai tiếp 3 dây.")
    c_tot_rem.font = bold_font
    c_tot_rem.alignment = left_align
    c_tot_rem.border = total_border

    # Section 2: Detailed Line Inventory
    ws1.cell(row=16, column=1, value="II. DANH SÁCH CHI TIẾT THEO TỪNG LINE VÀ THIẾT BỊ (라인별 상세 내역)").font = section_font
    
    detail_headers = ["STT", "Phân Khu (구역)", "Mã Line", "Tên Thiết Bị / Vị Trí Trên Layout", "Dây Cần", "Đã Kéo", "Còn Thiếu", "Trạng Thái", "Độ Khả Thi", "Ghi Chú Kỹ Thuật (기술 메모)"]
    ws1.row_dimensions[17].height = 24
    for col_idx, h in enumerate(detail_headers, start=1):
        c = ws1.cell(row=17, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = cell_border
        
    detail_data = [
        (1, "Aging", "Line 1", "35105 - Aging #1", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây"),
        (2, "Aging", "Line 2", "35105 - Aging #2", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây"),
        (3, "Aging", "Line 3", "35105 - Aging #3", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây"),
        (4, "Aging", "Line 4", "35105 - Aging #4", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây"),
        (5, "Aging", "Line 5", "35105 - Aging #5", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây"),
        (6, "Aging", "Line 6", "2245 - Aging #1 (Khoanh đỏ)", 1, 0, 1, "Cần kéo thêm", "Có thể kéo", "Thiếu 1 dây, hạ tầng đã sẵn sàng kéo"),
        (7, "Winding", "Line 1", "Máy Quấn Winding Line 1", 2, 2, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 2/2 dây"),
        (8, "Winding", "Line 2", "Máy Quấn Winding Line 2", 6, 6, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 6/6 dây"),
        (9, "Winding", "Line 3", "Máy Quấn Winding Line 3", 4, 2, 2, "Cần kéo thêm", "Có thể kéo", "Hiện đã có 2/4 dây, cần kéo thêm 2 dây"),
        (10, "Winding", "Line 4", "Máy Quấn Winding Line 4", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (11, "Winding", "Line 5", "Máy Quấn Winding Line 5", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (12, "Winding", "Line 6", "Máy Quấn Winding Line 6", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (13, "Winding", "Line 7", "Máy Quấn Winding Line 7", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (14, "Winding", "Line 8", "Máy Quấn Winding Line 8", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (15, "Winding", "Line 9", "Máy Quấn Winding Line 9", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "Chưa có máng cáp / Vị trí máy bị vướng"),
        (16, "Assembly", "Line 9", "35105 Cell Ass'y Cụm 1 (Cell #1-#3)", 5, 5, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 5/5 dây cho cụm máy"),
        (17, "Assembly", "Line 10", "35105 Cell Ass'y Cụm 2 (Cell #1-#3)", 5, 5, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 5/5 dây cho cụm máy"),
        (18, "Assembly", "Line 1-8", "Khu Vực Lắp Ráp Line 1 -> Line 8", "-", 0, "-", "Chưa thể kéo", "Nghẽn hạ tầng", "8 vị trí chưa có máng cáp / chưa thể kéo"),
        (19, "Phụ trợ", "Partleader", "Bàn Làm Việc Partleader / VPSX", 1, 1, 0, "Hoàn thành", "Đã xong", "1 Dây Trunk/Uplink -> Switch 8P (chia máy bàn)"),
        (20, "Phụ trợ", "X-Ray", "Máy Đo Kiểm Soi X-Ray", 1, 1, 0, "Hoàn thành", "Đã xong", "Đã kéo hoàn thiện 1/1 dây kết nối máy tính X-Ray"),
    ]
    
    for r_idx, row_values in enumerate(detail_data, start=18):
        ws1.row_dimensions[r_idx].height = 20
        status_txt = row_values[7]
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            
            if c_idx in [1, 2, 3, 5, 6, 7]:
                cell.alignment = center_align
            elif c_idx in [8, 9]:
                cell.alignment = center_align
                if status_txt == "Hoàn thành":
                    cell.fill = done_fill
                    cell.font = done_font
                elif status_txt == "Cần kéo thêm":
                    cell.fill = pending_fill
                    cell.font = pending_font
                else:
                    cell.fill = blocked_fill
                    cell.font = blocked_font
            else:
                cell.alignment = left_align

    col_widths_1 = {1: 6, 2: 12, 3: 10, 4: 35, 5: 10, 6: 10, 7: 11, 8: 15, 9: 14, 10: 45}
    for col_idx, width in col_widths_1.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    # =============================================================
    # SHEET 2: KIỂM KÊ THIẾT BỊ IT XƯỞNG (IT ASSET INVENTORY)
    # =============================================================
    ws2 = wb.create_sheet(title="2. KIEM KE THIET BI IT XUONG")
    ws2.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws2.merge_cells("A1:K1")
    ws2["A1"] = "VINATECH CO., LTD. — FACTORY IT ASSET & DEVICE INVENTORY"
    ws2["A1"].font = Font(name=font_family, size=10, bold=True, color="D9E1F2")
    ws2["A1"].fill = primary_fill
    ws2["A1"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.row_dimensions[1].height = 20

    ws2.merge_cells("A2:K2")
    ws2["A2"] = "DANH MỤC KIỂM KÊ THIẾT BỊ IT CELL LINE DƯỚI XƯỞNG (CELL 라인 IT 설비 실사 대장)"
    ws2["A2"].font = title_font
    ws2["A2"].fill = primary_fill
    ws2["A2"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.row_dimensions[2].height = 28

    ws2.merge_cells("A3:K3")
    ws2["A3"] = "Cập nhật thực tế: 07/10/2026 | Báo cáo: Mr. Hải / Mr. Lee Sang Hwan | IT Phụ trách: Nguyễn Văn Đức (vanduc)"
    ws2["A3"].font = subtitle_font
    ws2["A3"].fill = primary_fill
    ws2["A3"].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.row_dimensions[3].height = 18

    # KPI Summary Cards for IT Assets (Row 5 - 6)
    kpi_assets = [
        ("B5:C5", "B6:C6", "🖥️ TỔNG PC DESKTOP (데스크탑)", "6 Máy", "B5", "B6", primary_fill, Font(name=font_family, size=15, bold=True, color="1B365D")),
        ("D5:E5", "D6:E6", "💻 TỔNG LAPTOP (노트북)", "3 Máy", "D5", "D6", done_fill, Font(name=font_family, size=15, bold=True, color="375623")),
        ("F5:G5", "F6:G6", "📱 TỔNG TRẠM KIOSK POP (키오스크)", "6 Trạm", "F5", "F6", pending_fill, Font(name=font_family, size=15, bold=True, color="833C0C")),
        ("H5:I5", "H6:I6", "🔌 TỔNG THIẾT BỊ MÁY TÍNH (총 PC/IT)", "15 Thiết bị", "H5", "H6", PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"), Font(name=font_family, size=15, bold=True, color="1B365D")),
    ]
    
    ws2.row_dimensions[5].height = 18
    ws2.row_dimensions[6].height = 26
    
    for t_range, v_range, title, val, t_cell, v_cell, fill, f_font in kpi_assets:
        ws2.merge_cells(t_range)
        ws2[t_cell] = title
        ws2[t_cell].font = kpi_title_font
        ws2[t_cell].fill = gray_fill
        ws2[t_cell].alignment = center_align
        ws2[t_cell].border = cell_border
        
        ws2.merge_cells(v_range)
        ws2[v_cell] = val
        ws2[v_cell].font = f_font
        ws2[v_cell].fill = fill
        ws2[v_cell].alignment = center_align
        ws2[v_cell].border = cell_border

    # Section Header
    ws2.cell(row=8, column=1, value="DANH SÁCH THIẾT BỊ IT THỰC TẾ ĐANG VẬN HÀNH DƯỚI XƯỞNG (실사 내역)").font = section_font
    
    asset_headers = ["STT", "Phân Khu / Bộ Phận", "Vị Trí Cụ Thể", "Loại Thiết Bị (Device)", "Tên Thiết Bị / Mô Tả", "Số Lượng", "Loại Kết Nối Mạng", "Cổng Kết Nối / Switch", "Trạng Thái Hoạt Động", "Người Sử Dụng / Phụ Trách", "Ghi Chú Nghiệp Vụ"]
    ws2.row_dimensions[9].height = 24
    for col_idx, h in enumerate(asset_headers, start=1):
        c = ws2.cell(row=9, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = cell_border
        
    # User's exact inventory:
    # 2 PC + 2 Laptop ở khu vực partleader
    # 1 Laptop ở khu vực máy x-ray
    # 1 PC ở OQC
    # 1 PC ở ngoại quan đóng gói
    # 1 PC ở đóng gói
    # 1 PC ở thành phẩm
    # 4 Kiosk ở 4 line: 1, 2, 9, 10
    # 1 Kiosk ở aging
    # 1 Kiosk ở điện cực
    # Switch & Hạ tầng
    
    asset_data = [
        # KHU VỰC PARTLEADER
        (1, "Văn Phòng Xưởng", "Khu vực Partleader", "PC Desktop", "PC Văn Phòng Partleader #1", 1, "LAN (Dây mạng)", "Port 2 - Switch 8P Partleader", "🟢 Hoạt động tốt", "Partleader / Quản đốc SX", "Nhập liệu báo cáo & điều hành ca"),
        (2, "Văn Phòng Xưởng", "Khu vực Partleader", "PC Desktop", "PC Văn Phòng Partleader #2", 1, "LAN (Dây mạng)", "Port 3 - Switch 8P Partleader", "🟢 Hoạt động tốt", "Tổ phó / Thống kê SX", "Theo dõi kế hoạch & lệnh SX"),
        (3, "Văn Phòng Xưởng", "Khu vực Partleader", "Laptop", "Laptop Kỹ Thuật Partleader #1", 1, "LAN / Wi-Fi", "Port 4 - Switch 8P Partleader", "🟢 Hoạt động tốt", "Kỹ sư Sản Xuất (PE)", "Giám sát quy trình kỹ thuật"),
        (4, "Văn Phòng Xưởng", "Khu vực Partleader", "Laptop", "Laptop Kỹ Thuật Partleader #2", 1, "LAN / Wi-Fi", "Port 5 - Switch 8P Partleader", "🟢 Hoạt động tốt", "Kỹ sư Thiết Bị (ME)", "Theo dõi thông số máy móc"),
        
        # KHU VỰC X-RAY
        (5, "Đo Kiểm Chất Lượng", "Khu vực Máy X-Ray", "Laptop", "Laptop Điều Khiển & Soi X-Ray", 1, "LAN (Dây mạng)", "Dây mạng X-Ray (1/1 dây)", "🟢 Hoạt động tốt", "OP / Kỹ thuật X-Ray", "Phân tích hình ảnh cấu trúc Cell"),
        
        # CÁC TRẠM QC & ĐÓNG GÓI
        (6, "Quản Lý Chất Lượng", "Khu vực OQC", "PC Desktop", "PC Trạm Kiểm Tra OQC", 1, "LAN (Dây mạng)", "Điểm mạng trạm OQC", "🟢 Hoạt động tốt", "Nhân viên QC (OQC)", "Kiểm tra chất lượng xuất xưởng"),
        (7, "Đóng Gói & Xuất Hàng", "Ngoại quan đóng gói", "PC Desktop", "PC Trạm Kiểm Tra Ngoại Quan", 1, "LAN (Dây mạng)", "Điểm mạng bàn Ngoại quan", "🟢 Hoạt động tốt", "OP Ngoại quan", "Kiểm tra ngoại quan trước đóng thùng"),
        (8, "Đóng Gói & Xuất Hàng", "Khu vực Đóng Gói", "PC Desktop", "PC Trạm Đóng Gói (Packing)", 1, "LAN (Dây mạng)", "Điểm mạng trạm Đóng gói", "🟢 Hoạt động tốt", "OP Đóng gói", "Đóng thùng, in tem Sanmina/thùng"),
        (9, "Kho & Thành Phẩm", "Khu vực Thành Phẩm", "PC Desktop", "PC Quản Lý Kho Thành Phẩm", 1, "LAN (Dây mạng)", "Điểm mạng kho Thành phẩm", "🟢 Hoạt động tốt", "Thủ kho Thành phẩm", "Nhập kho B540 & xuất hàng"),

        # HỆ THỐNG KIOSK POP
        (10, "Sản Xuất Winding", "Line 1 Winding", "Kiosk POP", "Kiosk Cảm Ứng POP Line 1", 1, "LAN (Dây mạng)", "Dây mạng Line 1 Winding (2 dây)", "🟢 Hoạt động tốt", "Công nhân Line 1", "Chốt sản lượng POP Kiosk"),
        (11, "Sản Xuất Winding", "Line 2 Winding", "Kiosk POP", "Kiosk Cảm Ứng POP Line 2", 1, "LAN (Dây mạng)", "Dây mạng Line 2 Winding (6 dây)", "🟢 Hoạt động tốt", "Công nhân Line 2", "Chốt sản lượng POP Kiosk"),
        (12, "Sản Xuất Assembly", "Line 9 Assembly", "Kiosk POP", "Kiosk Cảm Ứng POP Line 9 (Cell #1-#3)", 1, "LAN (Dây mạng)", "Cụm mạng Line 9 (5 dây)", "🟢 Hoạt động tốt", "Công nhân Line 9", "Chốt công đoạn & nạp NVL"),
        (13, "Sản Xuất Assembly", "Line 10 Assembly", "Kiosk POP", "Kiosk Cảm Ứng POP Line 10 (Cell #1-#3)", 1, "LAN (Dây mạng)", "Cụm mạng Line 10 (5 dây)", "🟢 Hoạt động tốt", "Công nhân Line 10", "Chốt công đoạn & nạp NVL"),
        (14, "Sản Xuất Aging", "Khu vực Aging", "Kiosk POP", "Kiosk Cảm Ứng POP Aging", 1, "LAN (Dây mạng)", "Cáp mạng trạm Aging", "🟢 Hoạt động tốt", "Công nhân Aging", "Quản lý mẻ nạp/xả điện tích"),
        (15, "Sản Xuất Điện Cực", "Khu vực Điện Cực", "Kiosk POP", "Kiosk Cảm Ứng POP Điện Cực", 1, "LAN (Dây mạng)", "Cáp mạng trạm Cắt/Trộn", "🟢 Hoạt động tốt", "Công nhân Điện cực", "Quản lý cuộn điện cực & Slitting"),

        # THIẾT BỊ HẠ TẦNG & NGOẠI VI
        (16, "Hạ Tầng Mạng", "Khu vực Partleader", "Switch Mạng", "Switch 8-Port Unmanaged", 1, "Trunk Uplink", "Dây Trunk từ tủ Switch chính", "🟢 Hoạt động tốt", "IT Đức Nguyễn", "Cấp mạng cho 2 PC + 2 Laptop"),
        (17, "Thiết Bị Ngoại Vi", "Đóng gói / Kiosk", "Máy In Mã Vạch", "Zebra / Citizen ZT411 (LAN/USB)", 2, "LAN / USB", "Kết nối PC Đóng gói & Kiosk", "🟢 Hoạt động tốt", "OP Sản xuất", "In tem Barcode thùng & Box"),
        (18, "Thiết Bị Ngoại Vi", "Kiosk POP & QC", "Máy Quét Mã Vạch", "Honeywell Xenon 1900GHD Barcode", 6, "USB Interface", "Gắn trực tiếp vào các Kiosk POP", "🟢 Hoạt động tốt", "OP Sản xuất", "Quét tem NVL, Lot, Box")
    ]
    
    for r_idx, row_values in enumerate(asset_data, start=10):
        ws2.row_dimensions[r_idx].height = 20
        status_txt = row_values[8]
        for c_idx, val in enumerate(row_values, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            
            if c_idx in [1, 6]:
                cell.alignment = center_align
            elif c_idx == 9:
                cell.alignment = center_align
                if "🟢" in status_txt:
                    cell.fill = done_fill
                    cell.font = done_font
                elif "🟡" in status_txt:
                    cell.fill = pending_fill
                    cell.font = pending_font
                else:
                    cell.fill = blocked_fill
                    cell.font = blocked_font
            else:
                cell.alignment = left_align

    # Total row for Sheet 2
    tot_row = len(asset_data) + 10
    ws2.row_dimensions[tot_row].height = 22
    ws2.cell(row=tot_row, column=1, value="").border = total_border
    c_tot_label2 = ws2.cell(row=tot_row, column=2, value="TỔNG THIẾT BỊ IT (합계)")
    c_tot_label2.font = bold_font
    c_tot_label2.alignment = left_align
    c_tot_label2.border = total_border
    
    ws2.merge_cells(f"C{tot_row}:E{tot_row}")
    c_tot_desc = ws2.cell(row=tot_row, column=3, value="6 PC Desktop + 3 Laptop + 6 Kiosk POP + 1 Switch 8P + 8 Thiết bị ngoại vi")
    c_tot_desc.font = bold_font
    c_tot_desc.alignment = left_align
    c_tot_desc.border = total_border
    
    c_tot_qty = ws2.cell(row=tot_row, column=6, value=f"=SUM(F10:F{tot_row-1})")
    c_tot_qty.font = bold_font
    c_tot_qty.alignment = center_align
    c_tot_qty.border = total_border
    
    for ci in range(7, 12):
        ws2.cell(row=tot_row, column=ci, value="").border = total_border

    # Auto-fit columns Sheet 2
    col_widths_2 = {1: 6, 2: 20, 3: 24, 4: 18, 5: 32, 6: 10, 7: 18, 8: 26, 9: 18, 10: 24, 11: 34}
    for col_idx, width in col_widths_2.items():
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    # Save files
    out_desktop = r"C:\Users\User Vinatech.DESKTOP-RJJSEQU\Desktop\BAO_CAO_TIEN_DO_MANG_VA_THIET_BI_IT_CELL_LINE.xlsx"
    out_mes = r"C:\Users\User Vinatech.DESKTOP-RJJSEQU\Desktop\PROCESS\MES_POP\doc\BAO_CAO_TIEN_DO_MANG_VA_THIET_BI_IT_CELL_LINE.xlsx"
    
    os.makedirs(os.path.dirname(out_mes), exist_ok=True)
    
    wb.save(out_desktop)
    wb.save(out_mes)
    print(f"Report created successfully at:\n- {out_desktop}\n- {out_mes}")

if __name__ == "__main__":
    create_excel_report()
